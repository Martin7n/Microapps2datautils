import os
import openpyxl

from csv_eea_co.eaa_data_handler import field_list
from esg_vehicles.main_esg import car_data, all_in_one, object_data_update
from esg_vehicles.weight_check import classify_by_weight, weight_normalization


def xls_parse_sample(file):
    wb = openpyxl.load_workbook(file)
    ws = wb.active
    interval = 10

    car_objects = []
    object_to_write = []

    for row in ws.iter_rows(min_row=2,
                            max_row=ws.max_row,
                            # max_row = interval,
                            min_col=1,
                            max_col=ws.max_column, values_only=True,
                            ):
            original_car_data = {
                "appendix": row[0],
                "vin": row[1],
                "brand": row[2],
                "model": row[3],
                "description": row[4],
                "weight": row[5],
                "measure_unit": row[6],
                "fuel": row[7],
                "category": row[8],
                "emission": row[9],
                "mileage": row[10],
                "reg_no": row[11]
            }

            car_objects.append(original_car_data)

    for single_obj in car_objects:
        updated_object =object_data_update(single_obj)
        object_to_write.append(updated_object)

    return object_to_write
        # for k,v in x.items():
        #     print(f"{k}:{v}", end="\t")







def xls_parse_base(file):

    wb = openpyxl.load_workbook(file)
    ws = wb.active
    interval = 10

    car_objects = {}

    for row in ws.iter_rows(min_row=2,
            max_row=ws.max_row,
            # max_row = interval,
            min_col=1,
            max_col=ws.max_column, values_only=True,
    ):
        brand = row[1]
        model = row[2]
        emission = row[3]
        fuel = row[5]
        category = row[7]
        if brand and model and emission and fuel and category:
            mark_model = f"{brand.lower()} {str(model).lower().replace('/', '_').strip()}"
            print(mark_model)
            if mark_model not in car_objects:
                car_objects[mark_model] = []
            car_objects[mark_model].append(fuel)

    # print(car_objects)
    for k,v in car_objects.items():
        print(k, v, len(v))



def data_collection(data):
    pass


def xls_writer(data):
    output_wb = openpyxl.Workbook()
    output_ws = output_wb.active
    output_ws.title = "Results"

    output_wb.save("output.xlsx")


if __name__ == '__main__':
    # xls_parse_base(r"C:\drob\ress.xlsx")
    xls_parse_sample(r"C:\drob\sample.xlsx")